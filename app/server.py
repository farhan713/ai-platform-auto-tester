"""
Flask web UI for the Skylar IQ QA tool — Postgres-backed, multi-user.

Run:
    python -m app.server
"""
from __future__ import annotations

import csv
import io
import json
import os
import platform
import queue
import re
import secrets
import sys
import threading
import time
import traceback
import uuid
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from flask import (
    Flask, render_template, request, jsonify, redirect, url_for,
    send_from_directory, Response, stream_with_context, abort, g, session, flash,
)

from app.runner import RunConfig, run as run_qa
from app.runner_sitesearch import run_sitesearch
from app.report import generate as generate_report
from app.excel_reader import read_questions
from app import db, auth


APP_VERSION = "3.0.0"

ROOT = Path(__file__).resolve().parent.parent
RUNS_DIR = ROOT / "runs"
DATA_DIR = ROOT / "data"
TEST_FILES_DIR = DATA_DIR / "test_files"
BUNDLES_DIR = DATA_DIR / "bundles"
RUNS_DIR.mkdir(exist_ok=True)
DATA_DIR.mkdir(exist_ok=True)
TEST_FILES_DIR.mkdir(exist_ok=True)
BUNDLES_DIR.mkdir(exist_ok=True)


app = Flask(
    __name__,
    template_folder=str(Path(__file__).parent / "templates"),
    static_folder=str(Path(__file__).parent / "static"),
)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
app.config["SECRET_KEY"] = os.environ.get(
    "SQA_SECRET_KEY",
    "dev-only-change-me-in-prod-" + secrets.token_hex(8),
)
app.permanent_session_lifetime = timedelta(days=14)


# ---------------------------------------------------------------------------
# Boot — make sure the schema is in place before serving
# ---------------------------------------------------------------------------
_db_ready = False
_db_lock = threading.Lock()


def _ensure_db_ready() -> None:
    global _db_ready
    if _db_ready: return
    with _db_lock:
        if _db_ready: return
        db.init_schema()
        _db_ready = True


@app.before_request
def _boot():
    if not _db_ready:
        try:
            _ensure_db_ready()
        except Exception as e:
            return Response(
                f"Database not ready yet: {e}\nRetry in a moment.\n",
                503, {"Retry-After": "5"},
            )
    auth.load_current_user()


@app.context_processor
def inject_globals():
    return {"app_version": APP_VERSION, "active": "", "g": g}


# ---------------------------------------------------------------------------
# Auth pages
# ---------------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("dashboard"))
    error = None
    email = ""
    if request.method == "POST":
        email = (request.form.get("email") or "").strip()
        password = request.form.get("password") or ""
        user = auth.authenticate(email, password)
        if user:
            auth.login_user(user)
            nxt = session.pop("next", None) or url_for("dashboard")
            return redirect(nxt)
        error = "Email or password is incorrect."
    return render_template("login.html", error=error, email=email,
                           signup_allowed=auth.signup_allowed())


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if g.user:
        return redirect(url_for("dashboard"))
    if not auth.signup_allowed():
        return render_template("login.html",
                               error="Public signup is disabled. Ask your admin to create an account for you.",
                               signup_allowed=False)
    error = None
    name = email = ""
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        email = (request.form.get("email") or "").strip()
        password = request.form.get("password") or ""
        try:
            user = auth.create_user(email, password, name)
            auth.login_user(user)
            return redirect(url_for("dashboard"))
        except ValueError as e:
            error = str(e)
        except Exception as e:
            error = f"Signup failed: {e}"
    return render_template("signup.html", error=error, name=name, email=email)


@app.route("/logout")
def logout():
    auth.logout_user()
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _job_dir(run_id: str) -> Path:
    return RUNS_DIR / run_id


# Per-job in-memory event queues + stop flags
_event_queues: dict[str, queue.Queue[str]] = {}
_stop_events: dict[str, threading.Event] = {}

# Global concurrency cap. Each active run holds an open Playwright Chromium
# (~600 MB - 1 GB resident with the SQL Agent SPA loaded) plus the Python
# runner thread. Without a cap, N users clicking "Start" within the same
# minute all spawn Chromium simultaneously and OOM the container — every
# in-progress run then dies. Excess runs sit in 'queued' status (already
# the DB default) until a slot opens.
#
# Default of 2 fits comfortably in a 4 GiB / 2 vCPU container; bump
# SQA_MAX_CONCURRENT_RUNS if you size the container up. Setting it to 1
# fully serialises runs (safest for tight memory).
MAX_CONCURRENT_RUNS = max(1, int(os.environ.get("SQA_MAX_CONCURRENT_RUNS", "2")))
_run_slot_semaphore = threading.BoundedSemaphore(MAX_CONCURRENT_RUNS)


# ---------------------------------------------------------------------------
# Hosted bundle helpers — uploaded ssLibrary builds for users without a public
# search page. We unpack the upload into data/bundles/<id>/ and serve it via
# a synthesised wrapper page at /hosted/<id>/index.html.
# ---------------------------------------------------------------------------
def _save_bundle(uploaded, original_name: str, friendly_name: str,
                 user_id: str) -> dict[str, Any]:
    """Persist an uploaded bundle (.js or .zip) and return its metadata row."""
    import zipfile
    bundle_id = uuid.uuid4().hex  # 32 chars — unguessable URL slug
    bundle_dir = BUNDLES_DIR / bundle_id
    bundle_dir.mkdir(parents=True, exist_ok=False)

    safe_name = (original_name or "bundle").replace("/", "_").replace("\\", "_")
    lower = safe_name.lower()

    main_js: str | None = None
    main_css: str | None = None
    file_count = 0

    if lower.endswith(".zip"):
        # Save zip then extract to the bundle dir
        zip_path = bundle_dir / "_upload.zip"
        uploaded.save(str(zip_path))
        try:
            with zipfile.ZipFile(zip_path, "r") as zf:
                for member in zf.namelist():
                    if member.endswith("/"): continue       # skip dirs
                    if ".." in member: continue             # zip-slip safety
                    # Skip macOS / Windows metadata cruft that gets auto-added
                    # to zips and confuses the main-file picker below.
                    if "__MACOSX" in member.split("/"): continue
                    base = member.rsplit("/", 1)[-1]
                    if base.startswith("._"): continue       # AppleDouble resource forks
                    if base in (".DS_Store", "Thumbs.db"): continue
                    parts = member.split("/")
                    rel = "/".join(parts[1:]) if len(parts) > 1 and parts[0].endswith("dist") else member
                    rel = rel.replace("..", "_")
                    target = bundle_dir / rel
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with zf.open(member) as src, open(target, "wb") as dst:
                        dst.write(src.read())
                    file_count += 1
        finally:
            try: zip_path.unlink()
            except Exception: pass

        # Find the main JS + CSS — prefer files with "bundle" in the name,
        # otherwise the first .js / .css. Filter out macOS/Windows metadata
        # paths even if they slipped through.
        def _real_files(suffix: str) -> list:
            return sorted(
                p for p in bundle_dir.rglob(f"*{suffix}")
                if p.is_file()
                and "__MACOSX" not in p.parts
                and not p.name.startswith("._")
                and p.name not in (".DS_Store", "Thumbs.db")
            )
        js_files = _real_files(".js")
        css_files = _real_files(".css")
        if js_files:
            preferred = [p for p in js_files if "bundle" in p.name.lower()]
            chosen = preferred[0] if preferred else js_files[0]
            main_js = str(chosen.relative_to(bundle_dir))
        if css_files:
            preferred = [p for p in css_files if any(t in p.name.lower() for t in ("library", "bundle", "ss-"))]
            chosen = preferred[0] if preferred else css_files[0]
            main_css = str(chosen.relative_to(bundle_dir))
    elif lower.endswith(".js"):
        # Single JS file
        target = bundle_dir / "bundle.js"
        uploaded.save(str(target))
        main_js = "bundle.js"
        file_count = 1
    else:
        # Unknown — clean up + reject
        import shutil
        shutil.rmtree(bundle_dir, ignore_errors=True)
        raise ValueError("Bundle upload must be a .js file or .zip archive.")

    if not main_js:
        import shutil
        shutil.rmtree(bundle_dir, ignore_errors=True)
        raise ValueError("Could not find any .js file in the upload.")

    size = sum(p.stat().st_size for p in bundle_dir.rglob("*") if p.is_file())
    db.execute(
        """INSERT INTO hosted_bundles (id, user_id, name, original_filename,
                                        file_count, size_bytes, main_js, main_css)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""",
        (bundle_id, user_id, friendly_name[:80] or safe_name,
         safe_name, file_count, size, main_js, main_css),
    )
    return db.fetch_one("SELECT * FROM hosted_bundles WHERE id = %s", (bundle_id,))


def _user_bundle(bundle_id: str) -> dict[str, Any] | None:
    return db.fetch_one(
        "SELECT * FROM hosted_bundles WHERE id = %s AND user_id = %s",
        (bundle_id, auth.current_user_id()),
    )


def _delete_bundle(bundle_id: str, user_id: str) -> bool:
    row = db.fetch_one(
        "SELECT id FROM hosted_bundles WHERE id = %s AND user_id = %s",
        (bundle_id, user_id),
    )
    if not row: return False
    db.execute("DELETE FROM hosted_bundles WHERE id = %s", (bundle_id,))
    import shutil
    shutil.rmtree(BUNDLES_DIR / bundle_id, ignore_errors=True)
    return True


def _render_hosted_index(bundle: dict[str, Any]) -> str:
    """Synthesise the wrapper HTML that loads the user's bundle and calls
    ssLibrary.init() with placeholder values. The runner's add_init_script
    override substitutes those placeholders with the user's Tenant Override
    config at runtime."""
    main_js = bundle.get("main_js") or "bundle.js"
    main_css = bundle.get("main_css")
    css_link = f'<link rel="stylesheet" href="{main_css}">' if main_css else ""
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Skylar QA — Hosted test page</title>
{css_link}
<style>
  body {{ margin: 24px; font-family: -apple-system, BlinkMacSystemFont, sans-serif; }}
  .qa-banner {{
    position: fixed; top: 0; right: 0; padding: 6px 14px;
    background: rgba(26, 138, 58, 0.9); color: #fff; font-size: 11px;
    z-index: 9999; border-bottom-left-radius: 6px; pointer-events: none;
  }}
  #searchbox.inputWrap, #searchbox {{
    width: 320px; padding: 10px 14px; font-size: 14px;
    border: 2px solid #1a3a6e; border-radius: 4px; outline: none;
  }}
</style>
</head>
<body class="bodyWrap">
  <div class="qa-banner">Skylar QA hosted test page</div>
  <div class="mainwrap">
    <div>
      <input class="inputWrap" type="text" id="searchbox">
    </div>
  </div>
  <script src="{main_js}"></script>
  <script>
    // Placeholder init() call — these args are replaced at runtime by the
    // QA tool's Tenant Override config (org_id, console_url, jwt_user, etc.)
    if (window.ssLibrary && typeof window.ssLibrary.init === 'function') {{
      window.ssLibrary.init('placeholder', '', '', '', '', 'placeholder', 'placeholder');
    }} else {{
      console.error('[skylar-qa] window.ssLibrary not defined — bundle did not register the widget');
      document.body.insertAdjacentHTML('beforeend',
        '<div style="color:#c62828;margin:14px 0">Bundle did not register window.ssLibrary. Check the JS file.</div>');
    }}
  </script>
</body>
</html>"""


# ---------------------------------------------------------------------------
# Hosted-bundle API + serving routes
# ---------------------------------------------------------------------------
@app.route("/api/hosted-bundles", methods=["GET", "POST"])
@auth.login_required
def api_hosted_bundles():
    if request.method == "GET":
        rows = db.fetch_all(
            "SELECT * FROM hosted_bundles WHERE user_id = %s ORDER BY uploaded_at DESC",
            (auth.current_user_id(),))
        return jsonify({"bundles": [
            {"id": r["id"], "name": r["name"],
             "original_filename": r["original_filename"],
             "file_count": r["file_count"], "size_bytes": r["size_bytes"],
             "main_js": r["main_js"], "main_css": r["main_css"],
             "uploaded_at": r["uploaded_at"].isoformat() if r["uploaded_at"] else None,
             "test_url": f"/hosted/{r['id']}/index.html"}
            for r in rows]})
    f = request.files.get("file")
    if not f: return jsonify({"error": "no file uploaded (field 'file')"}), 400
    name = (request.form.get("name") or "").strip() or (f.filename or "bundle")
    try:
        meta = _save_bundle(f, f.filename or "bundle", name, auth.current_user_id())
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"upload failed: {e}"}), 500
    return jsonify({"bundle": {
        "id": meta["id"], "name": meta["name"],
        "original_filename": meta["original_filename"],
        "file_count": meta["file_count"], "size_bytes": meta["size_bytes"],
        "main_js": meta["main_js"], "main_css": meta["main_css"],
        "uploaded_at": meta["uploaded_at"].isoformat() if meta["uploaded_at"] else None,
        "test_url": f"/hosted/{meta['id']}/index.html",
    }}), 201


@app.route("/api/hosted-bundles/<bundle_id>", methods=["DELETE"])
@auth.login_required
def api_hosted_bundle_delete(bundle_id: str):
    if _delete_bundle(bundle_id, auth.current_user_id()):
        return jsonify({"ok": True})
    return jsonify({"error": "not found"}), 404


# Hosted serving — these routes are NOT login-protected (the runner's headless
# Chromium has no session cookie). Bundle IDs are 32-char random hex so URLs
# are unguessable; only the owner ever sees the URL via the API list.
@app.route("/hosted/<bundle_id>/")
@app.route("/hosted/<bundle_id>/index.html")
def hosted_index(bundle_id: str):
    row = db.fetch_one("SELECT * FROM hosted_bundles WHERE id = %s", (bundle_id,))
    if not row: abort(404)
    return Response(_render_hosted_index(row), mimetype="text/html; charset=utf-8")


@app.route("/hosted/<bundle_id>/<path:rel>")
def hosted_asset(bundle_id: str, rel: str):
    if not db.fetch_one("SELECT id FROM hosted_bundles WHERE id = %s", (bundle_id,)):
        abort(404)
    bundle_dir = BUNDLES_DIR / bundle_id
    full = (bundle_dir / rel).resolve()
    if not str(full).startswith(str(bundle_dir.resolve())): abort(403)
    if not full.exists(): abort(404)
    return send_from_directory(bundle_dir, rel)


def _emit(run_id: str, line: str) -> None:
    q = _event_queues.get(run_id)
    if q is not None:
        try: q.put_nowait(line)
        except queue.Full: pass
    log = _job_dir(run_id) / "events.log"
    try:
        with log.open("a") as f: f.write(line + "\n")
    except Exception: pass


def _run_owned_by_current_user(run_id: str) -> dict[str, Any] | None:
    """Return the run row only if owned by the current user."""
    return db.fetch_one(
        "SELECT * FROM runs WHERE id = %s AND user_id = %s",
        (run_id, auth.current_user_id()),
    )


def _all_results_for_run(run_id: str) -> list[dict[str, Any]]:
    rows = db.fetch_all(
        "SELECT record FROM query_results WHERE run_id = %s ORDER BY qid", (run_id,))
    return [r["record"] for r in rows]


# Warning patterns that are observations about the upstream LLM/SQL output,
# not real test failures. v3.7.7 demoted these from warn_reasons → info_reasons
# in the validator. Records saved BEFORE v3.7.7 still have them in warn_reasons,
# which keeps showing them as PARTIAL. We re-classify them here on display so
# the user sees consistent results across the whole history.
_INFO_GRADE_WARN_PATTERNS = (
    "empty/missing name",                                  # LLM forgot AS alias
    "columns differ between run-sql and generate-viz",     # alias vs raw expression
    "returned no rows",                                    # empty result set is a valid outcome
)


def _reclassify_legacy_validations(v: dict[str, Any], stored_status: str) -> tuple[str, list[str], list[str], list[str]]:
    """Re-apply v3.7.7 + v3.7.9 validation rules to a stored record.

    Returns (overall_status, fail_reasons, warn_reasons, info_reasons).

    Two corrections applied on top of older validators:

    1. Demoted warnings — "empty column name" / "columns differ between
       run-sql and generate-viz" are observations about the LLM-generated
       SQL, not test failures. Moved warn → info so they don't keep the
       record showing PARTIAL.

    2. Chain-of-blame fix — "generate-viz call missing" is a misleading
       fail when there's a clear upstream cause (run-sql timed out and
       submit_query bailed before checking the viz button, generate-sql
       returned 4xx, run-sql returned 0 rows). Strip the symptom and
       surface the real cause.
    """
    fails = list(v.get("fail_reasons") or [])
    warns_raw = list(v.get("warn_reasons") or [])
    infos = list(v.get("info_reasons") or [])

    # ---- step 1: demote info-grade warnings ----
    promoted_warns: list[str] = []
    real_warns: list[str] = []
    for w in warns_raw:
        if any(p in w for p in _INFO_GRADE_WARN_PATTERNS):
            promoted_warns.append(w)
        else:
            real_warns.append(w)
    for w in promoted_warns:
        if w not in infos:
            infos.append(w)

    # ---- step 2: chain-of-blame fix for "generate-viz call missing" ----
    timed_out_on = v.get("timed_out_on")
    gs_status = ((v.get("generate_sql_call") if False else None)
                 or v.get("generate_sql_status_legacy"))  # placeholder
    upstream_gen_sql_4xx = any(
        "generate-sql HTTP 4" in f or "generate-sql HTTP 5" in f for f in fails
    )
    run_sql_timed_out = bool(timed_out_on == "run-sql")
    run_sql_returned_empty = (v.get("run_sql_present") is True
                              and v.get("run_sql_has_rows") is False)
    viz_blocked_by_upstream = (
        upstream_gen_sql_4xx or run_sql_timed_out or run_sql_returned_empty
    )
    if viz_blocked_by_upstream:
        fails = [f for f in fails if f != "generate-viz call missing"]

    # ---- step 3: add a primary timeout reason if missing ----
    if stored_status == "TIMEOUT" and timed_out_on:
        timeout_msg = (
            f"{timed_out_on} did not respond within the configured timeout — "
            "downstream steps (visualization, validation) were skipped"
        )
        if not any(timeout_msg in w for w in real_warns):
            real_warns.insert(0, timeout_msg)

    # ---- step 4: compute the effective status ----
    # Re-derive status from scratch using the same rules as runner.validate_query.
    if fails:
        effective = "FAIL"
    elif stored_status == "TIMEOUT" or timed_out_on:
        effective = "TIMEOUT"
    elif real_warns:
        effective = "PARTIAL"
    else:
        effective = "PASS"
    return effective, fails, real_warns, infos


def _all_results_summary_for_run(run_id: str) -> list[dict[str, Any]]:
    """Slim per-query summary used by the run dashboard tabs.

    The full ``record`` JSONB carries every captured XHR/fetch (headers +
    full response bodies, including multi-MB run-sql payloads). For runs
    with hundreds of queries that adds up to hundreds of MB of JSON shipped
    on every dashboard refresh, which freezes the browser before it can
    even show the list. The dashboard only needs status, timing, validation
    reasons, and a per-call classification/status/duration count for the
    Network tab — so we project out exactly those fields here.

    Also applies v3.7.7's relaxed validation rules to records saved by older
    validators (see _reclassify_legacy_validations).
    """
    rows = db.fetch_all(
        "SELECT record FROM query_results WHERE run_id = %s ORDER BY qid", (run_id,))
    out: list[dict[str, Any]] = []
    for row in rows:
        rec = row["record"] or {}
        v = rec.get("validations") or {}
        rs = rec.get("run_sql_call") or {}
        gv = rec.get("generate_viz_call") or {}
        gs = rec.get("generate_sql_call") or {}

        slim_calls = []
        for c in (rec.get("calls") or []):
            slim_calls.append({
                "classification": c.get("classification"),
                "status": c.get("status"),
                "duration_ms": c.get("duration_ms"),
            })

        effective_status, fails, warns, infos = _reclassify_legacy_validations(
            v, rec.get("overall_status") or ""
        )

        out.append({
            "id": rec.get("id"),
            "nl_query": rec.get("nl_query"),
            "overall_status": effective_status,
            "stored_status": rec.get("overall_status"),  # original from DB, for debugging
            "total_duration_ms": rec.get("total_duration_ms"),
            "validations": {
                "fail_reasons": fails,
                "warn_reasons": warns,
                "info_reasons": infos,
                "run_sql_row_count": v.get("run_sql_row_count"),
                "run_sql_returned_sql": v.get("run_sql_returned_sql"),
            },
            "run_sql_call": {"status": rs.get("status")} if rs else None,
            "generate_viz_call": {"status": gv.get("status")} if gv else None,
            "generate_sql_call": {"status": gs.get("status")} if gs else None,
            "calls": slim_calls,
        })
    return out


# ---------------------------------------------------------------------------
# Variant-consistency report
# ---------------------------------------------------------------------------
def _normalize_sql(sql: Any) -> str:
    """Normalize SQL for 'same query' comparison: strip trailing semicolons,
    collapse all whitespace to single spaces, lowercase. (User chose
    'normalized exact match' — literals are NOT blanked, so WHERE x=7 and
    WHERE x=9 are treated as different.)"""
    if not sql:
        return ""
    s = str(sql).strip()
    s = s.rstrip(";").strip()
    s = re.sub(r"\s+", " ", s)
    return s.lower()


def _extract_sql_deep(body: Any) -> str | None:
    """Find the generated SQL string anywhere in a Celerant response body.

    The SQL the agent produced lives in the GENERATE-SQL response at
    responseBody.data.generated_sql.sql — nested several levels deep, not at
    the top level, and on the generate-sql call (not run-sql, whose body is
    the executed result rows). This walks the structure to find it robustly.
    """
    if isinstance(body, dict):
        gs = body.get("generated_sql")
        if isinstance(gs, dict) and isinstance(gs.get("sql"), str) and gs["sql"].strip():
            return gs["sql"]
        for k in ("sql", "sql_query", "executed_sql", "query"):
            v = body.get(k)
            if isinstance(v, str) and "select" in v.lower():
                return v
        for v in body.values():
            r = _extract_sql_deep(v)
            if r:
                return r
    elif isinstance(body, list):
        for v in body:
            r = _extract_sql_deep(v)
            if r:
                return r
    return None


def _sql_structure(sql: Any) -> dict[str, Any] | None:
    """Extract a structural fingerprint of a SQL query for semantic comparison:
      - outputs: the set of output column aliases (AS [x] / AS x)
      - tables:  the set of tables referenced after FROM / JOIN
      - where:   the WHERE predicate, normalized (alias prefixes stripped)
      - groupby: the GROUP BY clause, normalized

    This deliberately ignores HOW each output column is computed (the SELECT
    expressions), so two queries that hit the same tables, apply the same
    filter, and return the same columns are considered the same query even
    if one title-cases a column and the other doesn't.
    """
    if not sql or "select" not in str(sql).lower():
        return None
    s = re.sub(r"\s+", " ", str(sql).strip().rstrip(";")).strip()
    low = s.lower()

    outputs: set[str] = set()
    for m in re.findall(r"\bas\s+\[([^\]]+)\]", s, re.I):       # AS [Brand]
        outputs.add(m.strip().lower())
    for m in re.findall(r"\bas\s+([a-z_][\w]*)", s, re.I):      # AS Brand
        outputs.add(m.strip().lower())

    tables = set(t.lower() for t in re.findall(r"\b(?:from|join)\s+([a-z_][\w\.]*)", low))

    def _strip_aliases(txt: str) -> str:
        # remove "x." table-alias prefixes so s.sku_id == sku_id == t.sku_id
        txt = re.sub(r"\b[a-z_][\w]*\.", "", txt)
        return re.sub(r"\s+", " ", txt).strip()

    where = ""
    wm = re.search(r"\bwhere\b(.*?)(?:\bgroup\s+by\b|\border\s+by\b|\bhaving\b|$)", low)
    if wm:
        where = _strip_aliases(wm.group(1).strip())

    groupby = ""
    gm = re.search(r"\bgroup\s+by\b(.*?)(?:\border\s+by\b|\bhaving\b|$)", low)
    if gm:
        groupby = _strip_aliases(gm.group(1).strip())

    return {"outputs": outputs, "tables": tables, "where": where, "groupby": groupby}


def _semantic_compare(orig_sql: Any, var_sql: Any) -> tuple[bool | None, str]:
    """Return (equivalent, reason). equivalent is None when we can't parse
    one of the queries (e.g. the original produced no SQL)."""
    a = _sql_structure(orig_sql)
    b = _sql_structure(var_sql)
    if a is None or b is None:
        return None, "could not parse SQL for one of the queries"
    diffs = []
    if a["tables"] != b["tables"]:
        diffs.append("different tables")
    if a["where"] != b["where"]:
        diffs.append("different WHERE filter")
    if a["outputs"] != b["outputs"]:
        diffs.append("different output columns")
    if a["groupby"] != b["groupby"]:
        diffs.append("different GROUP BY")
    if not diffs:
        return True, "same tables, filter, output columns & grouping — only column formatting differs"
    return False, "; ".join(diffs)


def _extract_question_understanding_deep(body: Any) -> str | None:
    """Find generated_sql.question_understanding anywhere in a response body.
    Its value is either the literal 'rag_semantic' (SQL retrieved from RAG /
    trained examples) or a natural-language sentence (LLM-generated)."""
    if isinstance(body, dict):
        gs = body.get("generated_sql")
        if isinstance(gs, dict) and isinstance(gs.get("question_understanding"), str):
            return gs["question_understanding"]
        if isinstance(body.get("question_understanding"), str):
            return body["question_understanding"]
        for v in body.values():
            r = _extract_question_understanding_deep(v)
            if r is not None:
                return r
    elif isinstance(body, list):
        for v in body:
            r = _extract_question_understanding_deep(v)
            if r is not None:
                return r
    return None


def _sql_source_from_qu(qu: Any) -> str | None:
    """Map question_understanding to a source label: 'rag' or 'llm'."""
    if not qu or not isinstance(qu, str):
        return None
    return "rag" if qu.strip().lower() == "rag_semantic" else "llm"


def _variant_report(run_id: str, variant_groups: dict[str, Any]) -> dict[str, Any]:
    """Build the per-group SQL-consistency report from stored query records."""
    rows = db.fetch_all(
        "SELECT qid, record FROM query_results WHERE run_id = %s", (run_id,))
    by_qid: dict[int, dict[str, Any]] = {r["qid"]: (r["record"] or {}) for r in rows}

    def sql_of(qid: int):
        rec = by_qid.get(qid) or {}
        # 1) validator-populated field (newer runs)
        s = (rec.get("validations") or {}).get("run_sql_returned_sql")
        if s:
            return s
        # 2) deep-extract from the generate-sql call (then run-sql) for older
        #    records whose validator didn't capture the nested SQL path.
        for key in ("generate_sql_call", "run_sql_call"):
            s = _extract_sql_deep((rec.get(key) or {}).get("response_body"))
            if s:
                return s
        return None

    def understanding_of(qid: int) -> str | None:
        rec = by_qid.get(qid) or {}
        v = rec.get("validations") or {}
        if v.get("question_understanding") is not None:
            return v.get("question_understanding")
        return _extract_question_understanding_deep(
            (rec.get("generate_sql_call") or {}).get("response_body"))

    def source_of(qid: int) -> str | None:
        rec = by_qid.get(qid) or {}
        v = rec.get("validations") or {}
        if v.get("sql_source"):
            return v.get("sql_source")
        return _sql_source_from_qu(understanding_of(qid))

    def status_of(qid: int):
        return (by_qid.get(qid) or {}).get("overall_status")

    out_groups: list[dict[str, Any]] = []
    total_variants = matched_variants = equivalent_variants = 0
    groups_fully_exact = groups_fully_equivalent = 0
    rag_count = llm_count = 0  # across all questions (originals + variants)

    # Hierarchical display numbers: original = "1", its variants = "1.1",
    # "1.2", ...  (group index is 1-based and contiguous regardless of the
    # stored group_id). The underlying qid is kept for drill-down links.
    def _tally_source(src: str | None) -> None:
        nonlocal rag_count, llm_count
        if src == "rag":
            rag_count += 1
        elif src == "llm":
            llm_count += 1

    for gidx, g in enumerate((variant_groups.get("groups") or []), start=1):
        oqid = g["original_qid"]
        osql = sql_of(oqid)
        onorm = _normalize_sql(osql)
        _tally_source(source_of(oqid))
        variants = []
        for vi, vqid in enumerate(g.get("variant_qids", []), start=1):
            vsql = sql_of(vqid)
            # Exact (literal) match — normalized whitespace/case.
            exact = bool(onorm) and (_normalize_sql(vsql) == onorm)
            # Semantic equivalence — same tables/filter/output columns.
            equivalent, reason = _semantic_compare(osql, vsql)
            # Exact match always implies equivalent.
            if exact:
                equivalent, reason = True, "identical SQL"
            total_variants += 1
            if exact:
                matched_variants += 1
            if equivalent:
                equivalent_variants += 1
            _tally_source(source_of(vqid))
            rec = by_qid.get(vqid) or {}
            variants.append({
                "label": f"{gidx}.{vi}",
                "qid": vqid,
                "nl_query": rec.get("nl_query"),
                "status": status_of(vqid),
                "sql": vsql,
                "matches_original": exact,
                "equivalent": equivalent,
                "equivalent_reason": reason,
                "source": source_of(vqid),
                "question_understanding": understanding_of(vqid),
            })
        n = len(g.get("variant_qids", []))
        nexact = sum(1 for v in variants if v["matches_original"])
        nequiv = sum(1 for v in variants if v["equivalent"])
        if n and nexact == n:
            groups_fully_exact += 1
        if n and nequiv == n:
            groups_fully_equivalent += 1
        out_groups.append({
            "label": str(gidx),
            "group_id": g["group_id"],
            "sheet": g.get("sheet"),
            "original_qid": oqid,
            "original_text": g.get("original_text"),
            "original_sql": osql,
            "original_status": status_of(oqid),
            "original_has_sql": bool(onorm),
            "original_source": source_of(oqid),
            "original_understanding": understanding_of(oqid),
            "variant_count": n,
            "matched": nexact,
            "equivalent_count": nequiv,
            "match_rate": (nexact / n) if n else None,
            "equivalent_rate": (nequiv / n) if n else None,
            "variants": variants,
        })

    return {
        "groups": out_groups,
        "totals": {
            "groups": len(out_groups),
            "groups_fully_consistent": groups_fully_exact,
            "groups_fully_equivalent": groups_fully_equivalent,
            "variants": total_variants,
            "matched_variants": matched_variants,
            "equivalent_variants": equivalent_variants,
            "overall_match_rate": (matched_variants / total_variants) if total_variants else None,
            "overall_equivalent_rate": (equivalent_variants / total_variants) if total_variants else None,
            "rag_count": rag_count,
            "llm_count": llm_count,
        },
    }


# ---------------------------------------------------------------------------
# Background runner
# ---------------------------------------------------------------------------
def _run_job(run_id: str, cfg: RunConfig) -> None:
    """Run the QA, persist per-query rows + events. Filesystem still holds
    screenshots/REPORT/etc; DB holds run metadata + per-query JSONB.

    Acquires a slot from the global semaphore before starting the actual
    run. Excess concurrent submissions sit here in 'queued' status until
    an active run finishes. The DB row was already inserted as 'queued'
    by the /jobs POST handler so it shows up in the dashboard immediately,
    just without a started_at timestamp until the slot opens.
    """
    # Surface "waiting for slot" feedback to the user — the run page is
    # already loaded and listening on /events when this thread starts.
    waiting_emitted = False
    if not _run_slot_semaphore.acquire(blocking=False):
        _emit(run_id, f"[queued] waiting for a free run slot "
                      f"(cap = {MAX_CONCURRENT_RUNS} concurrent runs per replica)")
        waiting_emitted = True
        # Honour stop while queued — a user can cancel before the run starts.
        ev = _stop_events.get(run_id)
        while True:
            if ev is not None and ev.is_set():
                _emit(run_id, "[queued] cancelled before slot acquired")
                db.execute(
                    "UPDATE runs SET status = 'stopped', finished_at = NOW() WHERE id = %s",
                    (run_id,),
                )
                _emit(run_id, "[__end__]")
                return
            if _run_slot_semaphore.acquire(timeout=2.0):
                break
    try:
        if waiting_emitted:
            _emit(run_id, "[queued] slot acquired — starting")
        db.execute(
            "UPDATE runs SET status = 'running', started_at = NOW() WHERE id = %s",
            (run_id,),
        )
        cfg.on_event = lambda line: _emit(run_id, line)
        cfg.should_stop = lambda: _stop_events.get(run_id, threading.Event()).is_set()

        # Wrap the runner's per-query save so we can ALSO write each row into Postgres.
        from app.runner import save_query_result as orig_save
        def db_aware_save(qr, output_dir):
            path = orig_save(qr, output_dir)
            try:
                from dataclasses import asdict
                d = asdict(qr)
                db.execute(
                    """INSERT INTO query_results
                       (run_id, qid, status, duration_ms, nl_query, record)
                       VALUES (%s, %s, %s, %s, %s, %s)
                       ON CONFLICT (run_id, qid) DO UPDATE SET
                         status = EXCLUDED.status,
                         duration_ms = EXCLUDED.duration_ms,
                         nl_query = EXCLUDED.nl_query,
                         record = EXCLUDED.record""",
                    (run_id, qr.id, qr.overall_status, qr.total_duration_ms,
                     qr.nl_query, db.jsonify(d)),
                )
            except Exception as e:
                _emit(run_id, f"[db] failed to persist q{qr.id:02d}: {e}")
            return path
        # Monkey-patch just for this run — both runners use save_query_result
        import app.runner as runner_mod
        import app.runner_sitesearch as sitesearch_mod
        runner_mod.save_query_result = db_aware_save
        sitesearch_mod.save_query_result = db_aware_save
        try:
            if cfg.test_type == "site_search":
                summary = run_sitesearch(cfg)
            else:
                summary = run_qa(cfg)
        finally:
            runner_mod.save_query_result = orig_save
            sitesearch_mod.save_query_result = orig_save

        try:
            outputs = generate_report(cfg.output_dir)
            summary["report_outputs"] = outputs
        except Exception as e:
            _emit(run_id, f"[report] generation failed: {e}")

        db.execute(
            "UPDATE runs SET status = %s, finished_at = NOW(), summary = %s WHERE id = %s",
            ("stopped" if summary.get("stopped_by_user") else "done",
             db.jsonify(summary), run_id),
        )
        _emit(run_id, f"[done] {summary}")
    except Exception as e:
        tb = traceback.format_exc()
        _emit(run_id, f"[fatal] {type(e).__name__}: {e}")
        _emit(run_id, tb)
        db.execute(
            "UPDATE runs SET status = 'failed', finished_at = NOW(), error = %s, traceback = %s WHERE id = %s",
            (str(e), tb, run_id),
        )
    finally:
        # Always release the slot, even on crash, so the next queued run
        # can pick it up. BoundedSemaphore guards against double-release.
        try:
            _run_slot_semaphore.release()
        except ValueError:
            pass
        _emit(run_id, "[__end__]")


# ---------------------------------------------------------------------------
# Pages (all login-required)
# ---------------------------------------------------------------------------
@app.route("/")
@auth.login_required
def dashboard():
    return render_template("dashboard.html", active="dashboard")


@app.route("/runs")
@auth.login_required
def runs_page():
    return render_template("runs.html", active="runs")


@app.route("/runs/new")
@auth.login_required
def new_run_page():
    return render_template("new_run.html", active="new_run")


@app.route("/test-files")
@auth.login_required
def test_files_page():
    return render_template("test_files.html", active="test_files")


@app.route("/test-pages")
@auth.login_required
def test_pages_page():
    return render_template("test_pages.html", active="test_pages")


# ---------------------------------------------------------------------------
# Training — proxy to Celerant's train_sql_examples_backoffice_validated API.
# Lets a QA user upload a CSV of validated SQL examples (or point at a server
# filepath) and POST it to the tenant's training endpoint, then see which
# rows were saved and which failed — all without leaving the tool or fighting
# CORS (the call is made server-side).
# ---------------------------------------------------------------------------
CELERANT_TRAIN_BASE = "https://celerantai.com/sql_agent/train_sql_examples_backoffice_validated"


@app.route("/training")
@auth.login_required
def training_page():
    return render_template("training.html", active="training",
                           default_console="https://celerantai.com")


@app.route("/training/submit", methods=["POST"])
@auth.login_required
def training_submit():
    """Proxy a training request to the Celerant SQL-Agent training endpoint.

    The Celerant endpoint executes every CSV row's sql_query against the
    tenant's run-sql endpoint and only persists rows whose response status
    is 'success'; failed rows come back in the response. We run this call
    server-side (long timeout, no browser CORS) and hand the JSON straight
    back to the page.
    """
    import requests  # lazy import — only needed for this feature

    database_id = (request.form.get("database_id") or "").strip().strip("/")
    if not database_id:
        return jsonify({"ok": False, "error": "Database ID is required."}), 400

    # Optional console origin override (default celerantai.com).
    console = (request.form.get("console_url") or "https://celerantai.com").strip().rstrip("/")
    base = f"{console}/sql_agent/train_sql_examples_backoffice_validated"
    url = f"{base}/{database_id}/"

    # Query params — only send the ones the user actually set.
    params: dict[str, str] = {}
    filepath = (request.form.get("sql_examples_filepath") or "").strip()
    if filepath:
        params["sql_examples_filepath"] = filepath
    for key in ("skip_display_clean", "skip_nlq_variants"):
        val = request.form.get(key)
        if val in ("true", "false"):
            params[key] = val

    # Optional CSV file upload (multipart). One of file/filepath is required.
    files = None
    parsed_questions: list[dict[str, str]] = []
    upload = request.files.get("sql_examples_file")
    if upload and upload.filename:
        csv_bytes = upload.read()
        files = {
            "sql_examples_file": (
                upload.filename,
                csv_bytes,
                upload.mimetype or "text/csv",
            )
        }
        # Parse the CSV into questions so the page can offer a follow-up
        # "test these same questions on the SQL Agent" run after training succeeds.
        try:
            text = csv_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                nl = (row.get("natural_language_query") or "").strip()
                sql = (row.get("sql_query") or "").strip()
                if nl:
                    parsed_questions.append(
                        {"natural_language_query": nl, "expected_sql": sql}
                    )
        except Exception:
            parsed_questions = []
    if not files and not filepath:
        return jsonify({
            "ok": False,
            "error": "Provide either a CSV file to upload, or a server-side "
                     "sql_examples_filepath. At least one is required.",
        }), 400

    # Optional bearer token — HTTPBearer is defined in the spec but not
    # enforced for this endpoint; pass it through if the tenant requires it.
    headers: dict[str, str] = {}
    token = (request.form.get("bearer_token") or "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"

    started = time.time()
    try:
        resp = requests.post(url, params=params, files=files, headers=headers,
                             timeout=600)
    except requests.exceptions.SSLError:
        # Fall back to no-verify if the tenant uses a self-signed cert.
        try:
            resp = requests.post(url, params=params, files=files, headers=headers,
                                 timeout=600, verify=False)
        except Exception as e:
            return jsonify({"ok": False, "url": url,
                            "error": f"Request failed (SSL): {type(e).__name__}: {e}"}), 502
    except Exception as e:
        return jsonify({"ok": False, "url": url,
                        "error": f"Request failed: {type(e).__name__}: {e}"}), 502
    elapsed_ms = int((time.time() - started) * 1000)

    try:
        body = resp.json()
    except Exception:
        body = (resp.text or "")[:100_000]

    return jsonify({
        "ok": 200 <= resp.status_code < 300,
        "status_code": resp.status_code,
        "url": url,
        "params": params,
        "elapsed_ms": elapsed_ms,
        "response": body,
        # Parsed from the uploaded CSV so the page can offer a follow-up
        # "test these on the SQL Agent" run without re-uploading.
        "uploaded_questions": parsed_questions,
    })


@app.route("/training/test-run", methods=["POST"])
@auth.login_required
def training_test_run():
    """Create a normal SQL Agent run using the questions that were just
    uploaded to the training endpoint. Lets the tester immediately verify
    that asking those same NL questions through the SQL Agent now returns
    the trained SQL.
    """
    login_url = (request.form.get("login_url") or "").strip()
    cred_username = (request.form.get("username") or "").strip()
    cred_password = (request.form.get("password") or "")
    machine_id = (request.form.get("machine_id") or "100").strip()
    sql_agent_path = (request.form.get("sql_agent_path") or
                      "/backoffice/mv-assets/index-modern.html#/listScreen/sqlagent").strip()
    try:
        run_sql_to = int(request.form.get("run_sql_timeout_ms") or 120_000)
        gen_viz_to = int(request.form.get("gen_viz_timeout_ms") or 120_000)
    except ValueError:
        return ("Timeout fields must be integers.", 400)

    if not login_url:
        return ("Missing login_url.", 400)
    if login_url.lower().startswith("file:"):
        return ("file:// URLs cannot be reached from inside the container.", 400)
    if not (cred_username and cred_password):
        return ("Username and password are required.", 400)

    qjson = request.form.get("questions_json") or "[]"
    try:
        raw = json.loads(qjson)
    except Exception:
        return ("questions_json must be a JSON array.", 400)
    questions: list[dict[str, Any]] = []
    for i, q in enumerate(raw):
        nl = (q.get("natural_language_query") if isinstance(q, dict) else "") or ""
        sql = (q.get("expected_sql") if isinstance(q, dict) else "") or ""
        nl = str(nl).strip()
        if nl:
            questions.append({
                "id": len(questions) + 1,
                "natural_language_query": nl,
                "expected_sql": str(sql).strip(),
            })
    if not questions:
        return ("No valid questions to run.", 400)

    run_id = uuid.uuid4().hex[:8] + "-" + datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    job = _job_dir(run_id)
    job.mkdir(parents=True, exist_ok=True)
    # Write an xlsx mirror of the question list so the run zip includes it
    # alongside everything else, matching what a normal run looks like.
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["natural_language_query", "expected_sql"])
        for q in questions:
            ws.append([q["natural_language_query"], q["expected_sql"]])
        wb.save(str(job / "questions.xlsx"))
    except Exception as e:
        return (f"Failed to write questions.xlsx: {e}", 500)

    db.execute(
        """INSERT INTO runs (id, user_id, test_type, login_url, username, machine_id,
                             sql_agent_path, search_input_selector, question_count,
                             test_file_name, status)
           VALUES (%s, %s, 'sql_agent', %s, %s, %s, %s, %s, %s, %s, 'queued')""",
        (run_id, auth.current_user_id(), login_url, cred_username, machine_id,
         sql_agent_path, "#searchbox", len(questions),
         "training-test-run.xlsx"),
    )
    cfg = RunConfig(
        login_url=login_url, username=cred_username, password=cred_password,
        questions=questions, output_dir=job, test_type="sql_agent",
        machine_id=machine_id, sql_agent_path=sql_agent_path,
        run_sql_timeout_ms=run_sql_to, gen_viz_timeout_ms=gen_viz_to,
        headless=True,
    )
    _event_queues[run_id] = queue.Queue(maxsize=10_000)
    _stop_events[run_id] = threading.Event()
    threading.Thread(target=_run_job, args=(run_id, cfg), daemon=True).start()
    return redirect(url_for("job_view", job_id=run_id))


# ---------------------------------------------------------------------------
# Assign Task — push a run's failed/partial/timeout queries to minijira
# (https://minijirabe-6fpw.onrender.com) as a task with the failed-question
# detail as the description. Credentials live in env vars, never in source.
# ---------------------------------------------------------------------------
MINIJIRA_BASE = os.environ.get("SQA_MINIJIRA_BASE_URL", "https://minijirabe-6fpw.onrender.com").rstrip("/")
_minijira_token: dict[str, Any] = {"value": None, "fetched_at": 0.0}
_minijira_token_lock = threading.Lock()


def _minijira_login_get_token(force: bool = False) -> str | None:
    """Log in to minijira and cache the JWT in-process. Refreshes on force
    or when the cached token is older than 30 minutes. Returns None if no
    credentials are configured (SQA_MINIJIRA_EMAIL / SQA_MINIJIRA_PASSWORD)."""
    import requests
    email = os.environ.get("SQA_MINIJIRA_EMAIL")
    password = os.environ.get("SQA_MINIJIRA_PASSWORD")
    if not (email and password):
        return None
    with _minijira_token_lock:
        if (not force) and _minijira_token["value"] and (
            time.time() - _minijira_token["fetched_at"] < 30 * 60):
            return _minijira_token["value"]
        resp = requests.post(f"{MINIJIRA_BASE}/api/v1/auth/login",
                             json={"email": email, "password": password},
                             timeout=30)
        resp.raise_for_status()
        token = (resp.json() or {}).get("access_token")
        if not token:
            return None
        _minijira_token["value"] = token
        _minijira_token["fetched_at"] = time.time()
        return token


def _minijira_get(path: str) -> Any:
    """GET against minijira with the cached token, refreshing on 401."""
    import requests
    token = _minijira_login_get_token()
    if not token:
        raise RuntimeError("minijira credentials not configured (set SQA_MINIJIRA_EMAIL + SQA_MINIJIRA_PASSWORD)")
    url = f"{MINIJIRA_BASE}{path}"
    resp = requests.get(url, headers={"Authorization": f"Bearer {token}"}, timeout=30)
    if resp.status_code == 401:
        token = _minijira_login_get_token(force=True)
        resp = requests.get(url, headers={"Authorization": f"Bearer {token}"}, timeout=30)
    resp.raise_for_status()
    return resp.json()


def _minijira_post(path: str, payload: dict[str, Any]) -> tuple[int, Any]:
    """POST against minijira with the cached token, refreshing on 401.
    Returns (status_code, parsed_body)."""
    import requests
    token = _minijira_login_get_token()
    if not token:
        raise RuntimeError("minijira credentials not configured (set SQA_MINIJIRA_EMAIL + SQA_MINIJIRA_PASSWORD)")
    url = f"{MINIJIRA_BASE}{path}"
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    resp = requests.post(url, headers=headers, json=payload, timeout=60)
    if resp.status_code == 401:
        token = _minijira_login_get_token(force=True)
        headers["Authorization"] = f"Bearer {token}"
        resp = requests.post(url, headers=headers, json=payload, timeout=60)
    try:
        body = resp.json()
    except Exception:
        body = (resp.text or "")[:5000]
    return resp.status_code, body


def _failed_questions_description(job_id: str) -> tuple[str, int]:
    """Build a human-readable list of FAIL / PARTIAL / TIMEOUT queries for
    this run, used as the (read-only) task description. Returns
    (text, failed_count)."""
    run = _run_owned_by_current_user(job_id)
    if not run:
        return "", 0
    rows = db.fetch_all(
        "SELECT qid, status, duration_ms, nl_query, record FROM query_results "
        "WHERE run_id = %s ORDER BY qid", (job_id,))
    bad: list[dict[str, Any]] = []
    for r in rows:
        rec = r.get("record") or {}
        v = rec.get("validations") or {}
        effective, fails, warns, _infos = _reclassify_legacy_validations(
            v, r["status"] or "")
        if effective in ("FAIL", "PARTIAL", "TIMEOUT"):
            reason = ""
            if fails: reason = fails[0]
            elif warns: reason = warns[0]
            bad.append({
                "qid": r["qid"], "nl": r["nl_query"] or "",
                "status": effective, "reason": reason,
                "duration_ms": r["duration_ms"],
            })
    # Full external URL to the run page so the dev opening the minijira task
    # can click straight through. url_for with _external uses the current
    # request's Host header — works behind the Container Apps proxy.
    try:
        run_url = url_for("job_view", job_id=job_id, _external=True)
    except Exception:
        run_url = ""
    lines = [
        f"Run: {job_id}",
    ]
    if run_url:
        lines.append(f"Run URL: {run_url}")
    lines.extend([
        f"Tenant: {run.get('login_url') or ''}",
        f"Total failing queries: {len(bad)}",
        "",
        "Failed / partial / timeout queries:",
        "",
    ])
    for i, q in enumerate(bad, start=1):
        dur = f" ({q['duration_ms']} ms)" if q['duration_ms'] is not None else ""
        lines.append(f"{i}. q{q['qid']:02d} [{q['status']}]{dur}")
        lines.append(f"   Question: {q['nl']}")
        if q['reason']:
            lines.append(f"   Reason:   {q['reason']}")
        lines.append("")
    return "\n".join(lines), len(bad)


@app.route("/jobs/<job_id>/assign-task")
@auth.login_required
def assign_task_page(job_id: str):
    run = _run_owned_by_current_user(job_id)
    if not run: abort(404)
    return render_template("assign_task.html",
                           job_id=job_id, active="runs",
                           creds_configured=bool(os.environ.get("SQA_MINIJIRA_EMAIL")
                                                 and os.environ.get("SQA_MINIJIRA_PASSWORD")))


@app.route("/jobs/<job_id>/assign-task/data")
@auth.login_required
def assign_task_data(job_id: str):
    """Populate the form: prefill description with failed-questions text, list
    mentionable users + projects so the dropdowns can be filled in one round-trip."""
    run = _run_owned_by_current_user(job_id)
    if not run: abort(404)
    description, failed_count = _failed_questions_description(job_id)
    try:
        users = _minijira_get("/api/v1/users/mentionable")
    except Exception as e:
        return jsonify({"ok": False, "error": f"users: {type(e).__name__}: {e}"}), 502
    try:
        projects = _minijira_get("/api/v1/projects/")
    except Exception as e:
        return jsonify({"ok": False, "error": f"projects: {type(e).__name__}: {e}"}), 502
    return jsonify({
        "ok": True,
        "title_suggestion": f"QA failures — Run {job_id}",
        "description": description,
        "failed_count": failed_count,
        "users": users,
        "projects": projects,
    })


@app.route("/jobs/<job_id>/assign-task/create", methods=["POST"])
@auth.login_required
def assign_task_create(job_id: str):
    """Create the task on minijira. Description is rebuilt server-side from
    the run record (the form's description field is read-only, so what the
    user 'sees' is what we send — but we never trust the client for that)."""
    run = _run_owned_by_current_user(job_id)
    if not run: abort(404)
    title = (request.form.get("title") or "").strip()
    priority = (request.form.get("priority") or "medium").strip().lower()
    if priority not in ("low", "medium", "high"):
        priority = "medium"
    assigned_to = (request.form.get("assignedTo") or "").strip()
    project_id = (request.form.get("projectId") or "").strip()
    due_date = (request.form.get("dueDate") or "").strip()  # YYYY-MM-DD
    tags_raw = (request.form.get("tags") or "").strip()
    tags = [t.strip() for t in re.split(r"[,\n]", tags_raw) if t.strip()]

    if not title:
        return jsonify({"ok": False, "error": "Title is required."}), 400
    if not assigned_to:
        return jsonify({"ok": False, "error": "Pick someone to assign to."}), 400
    if not project_id:
        return jsonify({"ok": False, "error": "Pick a project."}), 400
    if not due_date:
        return jsonify({"ok": False, "error": "Pick a due date."}), 400

    # The minijira API wants ISO datetime with timezone; the form gives us a
    # bare date. Use 18:30:00.000Z (matches the sample payload).
    iso_due = f"{due_date}T18:30:00.000Z" if re.match(r"^\d{4}-\d{2}-\d{2}$", due_date) else due_date

    # Description is ALWAYS the server-rebuilt failed-questions text, so the
    # user can't tamper with what gets attached.
    description, failed_count = _failed_questions_description(job_id)

    payload = {
        "title": title,
        "description": description,
        "priority": priority,
        "assignedTo": assigned_to,
        "projectId": project_id,
        "dueDate": iso_due,
        "tags": tags or ["qa-tool"],
    }
    try:
        status_code, body = _minijira_post("/api/v1/tasks/", payload)
    except Exception as e:
        return jsonify({"ok": False, "error": f"{type(e).__name__}: {e}"}), 502
    ok = 200 <= status_code < 300

    # Persist the assignment on the run so the Runs list can show the
    # assignee's initials and link back to the minijira task.
    if ok:
        assignee_name = (request.form.get("assignedToName") or "").strip()
        if not assignee_name:
            # Fall back to looking the user up from /users/mentionable in case
            # the page didn't forward the name (defence in depth).
            try:
                users = _minijira_get("/api/v1/users/mentionable")
                u = next((u for u in users if u.get("id") == assigned_to), None)
                if u:
                    assignee_name = u.get("name") or u.get("email") or ""
            except Exception:
                pass
        task_id = None
        if isinstance(body, dict):
            task_id = body.get("id") or body.get("_id") or body.get("taskId")
        try:
            db.execute(
                """UPDATE runs
                      SET assigned_to_name = %s,
                          assigned_to_id   = %s,
                          assigned_at      = NOW(),
                          minijira_task_id = %s
                    WHERE id = %s AND user_id = %s""",
                (assignee_name or None, assigned_to or None,
                 str(task_id) if task_id else None,
                 job_id, auth.current_user_id()),
            )
        except Exception:
            # Persistence is best-effort — don't block the success response.
            pass

    return jsonify({
        "ok": ok, "status_code": status_code,
        "task": body if ok else None,
        "error": None if ok else (body if isinstance(body, str) else (
            body.get("detail") or body.get("message") or str(body))[:500] if isinstance(body, dict) else "task create failed"),
        "failed_count": failed_count,
    })


# ---------------------------------------------------------------------------
# Activity Logs — pull /sql_agent/history_data/{org_id}/{offset}/{limit}/ for
# one or more tenants and present a per-org report (totals, LLM vs RAG split,
# status breakdown, searchable record list). The call is proxied server-side
# so the browser doesn't fight CORS and we can fan out across N orgs.
# ---------------------------------------------------------------------------
CELERANT_HISTORY_BASE = "https://celerantai.com/sql_agent/history_data"


def _yyyy_mm_dd_to_mm_dd_yyyy(s: str) -> str | None:
    """Convert HTML5 date input (YYYY-MM-DD) to the API's MM-DD-YYYY format."""
    if not s:
        return None
    s = s.strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if m:
        y, mo, d = m.groups()
        return f"{mo}-{d}-{y}"
    # Already MM-DD-YYYY?
    if re.match(r"^\d{2}-\d{2}-\d{4}$", s):
        return s
    return None


@app.route("/activity-logs")
@auth.login_required
def activity_logs_page():
    return render_template("activity_logs.html", active="activity_logs",
                           default_console="https://celerantai.com")


@app.route("/activity-logs/orgs")
@auth.login_required
def activity_logs_orgs():
    """Proxy to GET {console}/sql_agent/all_orgs/ so the dropdown can be
    populated server-side (no CORS, no Celerant token leak to the browser).
    Returns the list verbatim from responseBody.data, sorted by name."""
    import requests
    console = (request.args.get("console_url") or "https://celerantai.com").strip().rstrip("/")
    url = f"{console}/sql_agent/all_orgs/"
    headers: dict[str, str] = {}
    token = (request.args.get("bearer_token") or "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    try:
        resp = requests.get(url, headers=headers, timeout=30)
    except Exception as e:
        return jsonify({"ok": False, "error": f"{type(e).__name__}: {e}"}), 502
    if not (200 <= resp.status_code < 300):
        return jsonify({"ok": False, "status_code": resp.status_code,
                        "error": (resp.text or "")[:300]}), 502
    try:
        body = resp.json()
    except Exception:
        return jsonify({"ok": False, "error": "non-JSON response from all_orgs"}), 502
    orgs = (body.get("responseBody") or {}).get("data") or []
    # Project to the fields we actually use; sort by friendly name.
    slim = sorted([{
        "name": o.get("name") or "",
        "database_id": o.get("database_id") or "",
        "organization_id": o.get("organization_id") or "",
    } for o in orgs if o.get("database_id")],
        key=lambda o: (o["name"] or "").lower())
    return jsonify({"ok": True, "count": len(slim), "orgs": slim})


@app.route("/activity-logs/fetch", methods=["POST"])
@auth.login_required
def activity_logs_fetch():
    """Fetch history for each org id (one per line) and return a combined
    per-org report. No data is stored — this is a live proxy."""
    import requests

    raw_ids = (request.form.get("org_ids") or "").strip()
    org_ids = [s.strip() for s in re.split(r"[\s,]+", raw_ids) if s.strip()]
    if not org_ids:
        return jsonify({"ok": False, "error": "Provide at least one Organization ID."}), 400
    if len(org_ids) > 25:
        return jsonify({"ok": False, "error": "Up to 25 organizations per request."}), 400

    from_d = _yyyy_mm_dd_to_mm_dd_yyyy(request.form.get("from_date") or "")
    to_d   = _yyyy_mm_dd_to_mm_dd_yyyy(request.form.get("to_date") or "")
    if not from_d or not to_d:
        return jsonify({"ok": False, "error": "Provide both from_date and to_date (YYYY-MM-DD)."}), 400

    try:
        offset = max(0, int(request.form.get("offset") or 0))
        limit  = min(1000, max(1, int(request.form.get("limit") or 100)))
    except ValueError:
        return jsonify({"ok": False, "error": "offset and limit must be integers."}), 400

    console = (request.form.get("console_url") or "https://celerantai.com").strip().rstrip("/")
    base = f"{console}/sql_agent/history_data"
    headers: dict[str, str] = {}
    token = (request.form.get("bearer_token") or "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"

    def _stats(records: list[dict[str, Any]]) -> dict[str, Any]:
        llm = sum(1 for r in records if r.get("llm_invoked") is True)
        rag = sum(1 for r in records if r.get("llm_invoked") is False)
        statuses: dict[str, int] = {}
        users: set[str] = set()
        sessions: set[str] = set()
        first = last = None
        for r in records:
            s = (r.get("query_status") or "unknown").lower()
            statuses[s] = statuses.get(s, 0) + 1
            ln = r.get("login_name")
            if ln:
                users.add(str(ln))
            sid = r.get("session_id")
            if sid:
                sessions.add(str(sid))
            ts = r.get("created_at")
            if ts:
                first = ts if (first is None or ts < first) else first
                last  = ts if (last  is None or ts > last)  else last
        return {
            "llm": llm, "rag": rag,
            "complete": statuses.get("complete", 0),
            "failed":   statuses.get("failed", 0),
            "status_breakdown": statuses,
            "users":    sorted(users),
            "user_count": len(users),
            "session_count": len(sessions),
            "first_ts": first, "last_ts": last,
        }

    out_orgs: list[dict[str, Any]] = []
    grand = {"orgs": 0, "queries": 0, "llm": 0, "rag": 0, "complete": 0, "failed": 0}
    for org_id in org_ids:
        url = f"{base}/{org_id}/{offset}/{limit}/"
        try:
            resp = requests.get(url, params={"from_date": from_d, "to_date": to_d},
                                headers=headers, timeout=120)
            try:
                body = resp.json()
            except Exception:
                body = (resp.text or "")[:50_000]
            records = []
            history_count = None
            if isinstance(body, dict):
                rb = body.get("responseBody") or {}
                data = rb.get("data") if isinstance(rb, dict) else None
                if isinstance(data, dict):
                    records = data.get("history_records") or []
                    history_count = data.get("history_count")
            entry = {
                "org_id": org_id,
                "ok": 200 <= resp.status_code < 300,
                "status_code": resp.status_code,
                "url": url,
                "history_count": history_count,
                "count": len(records),
                "records": records,
                "stats": _stats(records),
                "error": None if 200 <= resp.status_code < 300 else (
                    body.get("responseHeader", {}).get("message") if isinstance(body, dict) else str(body)[:200]
                ),
            }
        except Exception as e:
            entry = {
                "org_id": org_id, "ok": False, "status_code": None,
                "url": url, "error": f"{type(e).__name__}: {e}",
                "count": 0, "records": [], "stats": _stats([]),
            }
        out_orgs.append(entry)
        if entry["ok"]:
            grand["orgs"] += 1
            grand["queries"] += entry["count"]
            grand["llm"] += entry["stats"]["llm"]
            grand["rag"] += entry["stats"]["rag"]
            grand["complete"] += entry["stats"]["complete"]
            grand["failed"] += entry["stats"]["failed"]

    return jsonify({
        "ok": True,
        "from_date": from_d, "to_date": to_d,
        "offset": offset, "limit": limit,
        "totals": grand,
        "organizations": out_orgs,
    })


# ---------------------------------------------------------------------------
# Variant Tests — run an original question + its paraphrases through the SQL
# AI Engine and report how many variants generate the same SQL as the original.
# Reuses the SQL Agent runner; the run is tagged with a variant_groups map.
# ---------------------------------------------------------------------------
@app.route("/variant-tests")
@auth.login_required
def variant_tests_page():
    return render_template("variant_tests.html", active="variant_tests")


@app.route("/variant-tests", methods=["POST"])
@auth.login_required
def variant_tests_create():
    from app.variant_reader import read_variant_questions

    login_url = (request.form.get("login_url") or "").strip()
    cred_username = (request.form.get("username") or "").strip()
    cred_password = (request.form.get("password") or "")
    machine_id = (request.form.get("machine_id") or "100").strip()
    sql_agent_path = (request.form.get("sql_agent_path") or
                      "/backoffice/mv-assets/index-modern.html#/listScreen/sqlagent").strip()
    run_sql_to = int(request.form.get("run_sql_timeout_ms") or 120_000)
    gen_viz_to = int(request.form.get("gen_viz_timeout_ms") or 120_000)
    f = request.files.get("questions")

    if not login_url:
        return ("Missing login_url", 400)
    if login_url.lower().startswith("file:"):
        return ("file:// URLs cannot be reached from inside the container. "
                "Use http:// or https://.", 400)
    if not (cred_username and cred_password):
        return ("Variant runs need a username and password", 400)
    if not f or not f.filename:
        return ("Upload a variant-questions .xlsx file", 400)

    run_id = uuid.uuid4().hex[:8] + "-" + datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    job = _job_dir(run_id)
    job.mkdir(parents=True, exist_ok=True)
    xlsx_path = job / "questions.xlsx"
    f.save(str(xlsx_path))

    try:
        questions, groups = read_variant_questions(xlsx_path)
    except Exception as e:
        return (f"Failed to read variant-questions file: {e}", 400)

    variant_groups = {"groups": groups}
    db.execute(
        """INSERT INTO runs (id, user_id, test_type, login_url, username, machine_id,
                             sql_agent_path, search_input_selector, question_count,
                             test_file_name, variant_groups, status)
           VALUES (%s, %s, 'sql_agent', %s, %s, %s, %s, %s, %s, %s, %s, 'queued')""",
        (run_id, auth.current_user_id(), login_url, cred_username, machine_id,
         sql_agent_path, "#searchbox", len(questions),
         (f.filename or "variant-questions.xlsx"),
         db.jsonify(variant_groups)),
    )

    cfg = RunConfig(
        login_url=login_url, username=cred_username, password=cred_password,
        questions=questions, output_dir=job,
        test_type="sql_agent",
        machine_id=machine_id, sql_agent_path=sql_agent_path,
        run_sql_timeout_ms=run_sql_to, gen_viz_timeout_ms=gen_viz_to,
        headless=True,
    )
    _event_queues[run_id] = queue.Queue(maxsize=10_000)
    _stop_events[run_id] = threading.Event()
    threading.Thread(target=_run_job, args=(run_id, cfg), daemon=True).start()
    return redirect(url_for("job_view", job_id=run_id))


@app.route("/jobs/<job_id>/variant")
@auth.login_required
def job_variant_view(job_id: str):
    run = _run_owned_by_current_user(job_id)
    if not run or not run.get("variant_groups"):
        abort(404)
    return render_template("variant_report.html", job_id=job_id, active="runs")


@app.route("/jobs/<job_id>/variant-report")
@auth.login_required
def job_variant_report_api(job_id: str):
    run = _run_owned_by_current_user(job_id)
    if not run:
        abort(404)
    vg = run.get("variant_groups")
    if not vg:
        return jsonify({"error": "This run is not a variant test."}), 400
    return jsonify(_variant_report(job_id, vg))


@app.route("/settings")
@auth.login_required
def settings_page():
    try:
        from playwright import __version__ as pw_version
    except Exception:
        pw_version = "unknown"
    return render_template(
        "settings.html",
        active="settings",
        python_version=sys.version.split()[0],
        playwright_version=pw_version,
        runs_dir=str(RUNS_DIR),
        presets_file="postgres://(database)",
    )


@app.route("/help")
@auth.login_required
def help_page():
    return render_template("help.html", active="help")


# ---------------------------------------------------------------------------
# Run creation + viewing (all owner-scoped)
# ---------------------------------------------------------------------------
@app.route("/jobs", methods=["POST"])
@auth.login_required
def jobs_create():
    test_type = (request.form.get("test_type") or "sql_agent").strip()
    if test_type not in ("sql_agent", "site_search"):
        return (f"Unknown test_type '{test_type}'", 400)

    login_url = (request.form.get("login_url") or "").strip()
    cred_username = (request.form.get("username") or "").strip()
    cred_password = (request.form.get("password") or "")
    machine_id = (request.form.get("machine_id") or "100").strip()
    sql_agent_path = (request.form.get("sql_agent_path") or
                      "/backoffice/mv-assets/index-modern.html#/listScreen/sqlagent").strip()
    search_input_selector = (request.form.get("search_input_selector") or "#searchbox").strip()
    run_sql_to = int(request.form.get("run_sql_timeout_ms") or 120_000)
    gen_viz_to = int(request.form.get("gen_viz_timeout_ms") or 120_000)
    test_file_id = (request.form.get("test_file_id") or "").strip()
    bundle_id = (request.form.get("bundle_id") or "").strip()
    f = request.files.get("questions")

    # Site Search runtime config — overrides the hardcoded ssLibrary.init() args
    # in the user's index.html so the same page can be tested against any tenant.
    site_search_config: dict[str, str] = {}
    if test_type == "site_search":
        for key in ("ss_org_id", "ss_console_url", "ss_server_url",
                    "ss_image_url", "ss_not_found_image_url",
                    "ss_jwt_user", "ss_jwt_pass"):
            v = (request.form.get(key) or "").strip()
            if v:
                site_search_config[key.replace("ss_", "")] = v
        if site_search_config and not site_search_config.get("server_url"):
            site_search_config["server_url"] = site_search_config.get("console_url", "")
        if site_search_config and not site_search_config.get("jwt_user"):
            site_search_config["jwt_user"] = site_search_config.get("org_id", "")

        # If a hosted bundle is picked, override login_url to point at our
        # generated wrapper page. The runner inside the container reaches it
        # via 127.0.0.1:5050 (same container).
        if bundle_id:
            owned = db.fetch_one(
                "SELECT id FROM hosted_bundles WHERE id = %s AND user_id = %s",
                (bundle_id, auth.current_user_id()),
            )
            if not owned:
                return ("Hosted bundle not found, or not yours.", 404)
            login_url = f"http://127.0.0.1:5050/hosted/{bundle_id}/index.html"

    if not login_url:
        return ("Missing login_url / target_url", 400)
    # Catch the file:// trap — Docker can't reach host filesystem paths.
    if login_url.lower().startswith("file:"):
        return (
            "file:// URLs cannot be reached from inside the Docker container.\n\n"
            "Use http:// or https:// instead. To test a local page, serve it over HTTP:\n"
            "    cd /path/to/your/page-folder\n"
            "    python3 -m http.server 8765\n"
            "Then use this URL:  http://host.docker.internal:8765/index.html",
            400,
        )
    if test_type == "sql_agent" and not (cred_username and cred_password):
        return ("SQL Agent runs need username and password", 400)
    if not f and not test_file_id:
        return ("Either upload a questions file or pick a saved test file", 400)

    run_id = uuid.uuid4().hex[:8] + "-" + datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    job = _job_dir(run_id)
    job.mkdir(parents=True, exist_ok=True)
    xlsx_path = job / "questions.xlsx"
    test_file_name = None

    if test_file_id:
        # Owner check: user can only use their own saved files
        tf = db.fetch_one(
            "SELECT * FROM test_files WHERE id = %s AND user_id = %s",
            (test_file_id, auth.current_user_id()),
        )
        if not tf:
            return ("Saved test file not found, or not yours.", 404)
        try:
            xlsx_path.write_bytes((TEST_FILES_DIR / tf["filename"]).read_bytes())
        except Exception as e:
            return (f"Failed to copy saved test file: {e}", 500)
        test_file_name = tf["original_name"]
    else:
        f.save(str(xlsx_path))

    try:
        questions = read_questions(xlsx_path)
    except Exception as e:
        return (f"Failed to read questions file: {e}", 400)

    db.execute(
        """INSERT INTO runs (id, user_id, test_type, login_url, username, machine_id,
                             sql_agent_path, search_input_selector, question_count,
                             test_file_id, test_file_name, site_search_config,
                             bundle_id, status)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'queued')""",
        (run_id, auth.current_user_id(), test_type, login_url, cred_username, machine_id,
         sql_agent_path, search_input_selector, len(questions),
         test_file_id or None, test_file_name,
         db.jsonify(site_search_config) if site_search_config else None,
         bundle_id or None),
    )

    cfg = RunConfig(
        login_url=login_url, username=cred_username, password=cred_password,
        questions=questions, output_dir=job,
        test_type=test_type,
        machine_id=machine_id, sql_agent_path=sql_agent_path,
        search_input_selector=search_input_selector,
        site_search_config=site_search_config or None,
        run_sql_timeout_ms=run_sql_to, gen_viz_timeout_ms=gen_viz_to,
        headless=True,
    )
    _event_queues[run_id] = queue.Queue(maxsize=10_000)
    _stop_events[run_id] = threading.Event()
    threading.Thread(target=_run_job, args=(run_id, cfg), daemon=True).start()
    return redirect(url_for("job_view", job_id=run_id))


@app.route("/jobs/<job_id>")
@auth.login_required
def job_view(job_id: str):
    run = _run_owned_by_current_user(job_id)
    if not run:
        abort(404)
    meta = {
        "test_type": run.get("test_type", "sql_agent"),
        "login_url": run["login_url"],
        "username": run["username"],
        "machine_id": run["machine_id"],
        "sql_agent_path": run["sql_agent_path"],
        "search_input_selector": run.get("search_input_selector"),
        "question_count": run["question_count"],
        "status": run["status"],
        "created_at": run["created_at"].isoformat() if run.get("created_at") else None,
        "started_at": run["started_at"].isoformat() if run.get("started_at") else None,
        "finished_at": run["finished_at"].isoformat() if run.get("finished_at") else None,
        "is_variant": bool(run.get("variant_groups")),
    }
    return render_template("job.html", job_id=job_id, meta=meta, active="runs")


@app.route("/jobs/<job_id>/meta")
@auth.login_required
def job_meta_api(job_id: str):
    """Lightweight JSON endpoint — used by query.html to know the test_type."""
    run = _run_owned_by_current_user(job_id)
    if not run: abort(404)
    return jsonify({"test_type": run.get("test_type", "sql_agent"),
                    "login_url": run["login_url"],
                    "question_count": run["question_count"]})


@app.route("/jobs/<job_id>/queries/<int:qid>")
@auth.login_required
def job_query_view(job_id: str, qid: int):
    if not _run_owned_by_current_user(job_id):
        abort(404)
    return render_template("query.html", job_id=job_id, qid=qid, active="runs")


@app.route("/jobs/<job_id>/status")
@auth.login_required
def job_status(job_id: str):
    run = _run_owned_by_current_user(job_id)
    if not run: abort(404)
    # Pull the full record JSONB for the re-validation pass — without it we'd
    # show stored PARTIALs that the v3.7.7 logic would consider PASS, which
    # would disagree with the All Queries table that DOES re-validate.
    rows = db.fetch_all(
        "SELECT qid, status, duration_ms, nl_query, record FROM query_results "
        "WHERE run_id = %s ORDER BY qid",
        (job_id,))
    counts = {"PASS": 0, "PARTIAL": 0, "FAIL": 0, "TIMEOUT": 0, "PENDING": 0}
    effective_by_qid: dict[int, str] = {}
    for r in rows:
        rec = r.get("record") or {}
        v = rec.get("validations") or {}
        effective, _, _, _ = _reclassify_legacy_validations(v, r["status"] or "")
        effective_by_qid[r["qid"]] = effective
        counts[effective] = counts.get(effective, 0) + 1
    last5 = [{"id": r["qid"], "nl_query": (r["nl_query"] or "")[:80],
              "status": effective_by_qid.get(r["qid"], r["status"]),
              "duration_ms": r["duration_ms"]} for r in rows[-5:]]
    stop_requested = bool(_stop_events.get(job_id) and _stop_events[job_id].is_set())
    return jsonify({
        "job_id": job_id, "status": run["status"],
        "stop_requested": stop_requested,
        "completed": sum(counts.values()) - counts["PENDING"],
        "total": run["question_count"],
        "counts": counts, "last5": last5,
    })


@app.route("/jobs/<job_id>/results")
@auth.login_required
def job_results(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    return jsonify({"results": _all_results_for_run(job_id)})


@app.route("/jobs/<job_id>/results-summary")
@auth.login_required
def job_results_summary(job_id: str):
    """Slim version of /results — used by the dashboard tabs (Overview,
    All queries, Network) to keep payload size bounded for large runs.
    Drops headers/bodies/etc; keeps the fields the UI actually renders."""
    if not _run_owned_by_current_user(job_id): abort(404)
    return jsonify({"results": _all_results_summary_for_run(job_id)})


@app.route("/jobs/<job_id>/events")
@auth.login_required
def job_events(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)

    @stream_with_context
    def gen():
        log = _job_dir(job_id) / "events.log"
        if log.exists():
            for line in log.read_text().splitlines():
                yield f"data: {line}\n\n"
        q = _event_queues.get(job_id)
        if q is None: return
        while True:
            try:
                line = q.get(timeout=20)
            except queue.Empty:
                yield ": keepalive\n\n"; continue
            yield f"data: {line}\n\n"
            if line == "[__end__]": return
    return Response(gen(), mimetype="text/event-stream")


@app.route("/jobs/<job_id>/stop", methods=["POST"])
@auth.login_required
def job_stop(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    ev = _stop_events.get(job_id)
    if ev is None:
        ev = threading.Event(); _stop_events[job_id] = ev
    ev.set()
    _emit(job_id, "[stop] user requested stop")
    return jsonify({"ok": True, "stop_requested": True})


@app.route("/jobs/<job_id>/result/<int:qid>")
@auth.login_required
def job_result(job_id: str, qid: int):
    if not _run_owned_by_current_user(job_id): abort(404)
    row = db.fetch_one(
        "SELECT record FROM query_results WHERE run_id = %s AND qid = %s",
        (job_id, qid),
    )
    if not row:
        return jsonify({"pending": True, "id": qid}), 202
    rec = row["record"] or {}
    # Apply the v3.7.7 re-validation so the detail page is consistent with
    # the All Queries table — otherwise an older record would still read
    # PARTIAL here even though the list view shows it as PASS.
    v = rec.get("validations") or {}
    effective, fails, warns, infos = _reclassify_legacy_validations(
        v, rec.get("overall_status") or ""
    )
    if rec.get("overall_status") != effective:
        rec["stored_status"] = rec.get("overall_status")
        rec["overall_status"] = effective
    v["fail_reasons"] = fails
    v["warn_reasons"] = warns
    v["info_reasons"] = infos
    rec["validations"] = v
    return jsonify(rec)


@app.route("/jobs/<job_id>/screenshots/<path:rel>")
@auth.login_required
def job_screenshot(job_id: str, rel: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    job = _job_dir(job_id)
    full = (job / "screenshots" / rel).resolve()
    if not str(full).startswith(str((job / "screenshots").resolve())): abort(403)
    if not full.exists(): abort(404)
    return send_from_directory(job / "screenshots", rel)


@app.route("/jobs/<job_id>/report")
@auth.login_required
def job_report(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    job = _job_dir(job_id)
    if not (job / "REPORT.html").exists():
        try: generate_report(job)
        except SystemExit: return ("No results yet", 404)
    return send_from_directory(job, "REPORT.html")


@app.route("/jobs/<job_id>/file/<path:rel>")
@auth.login_required
def job_file(job_id: str, rel: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    job = _job_dir(job_id)
    full = (job / rel).resolve()
    if not str(full).startswith(str(job.resolve())): abort(403)
    if not full.exists(): abort(404)
    return send_from_directory(job, rel)


@app.route("/jobs/<job_id>/export/csv")
@auth.login_required
def job_export_csv(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    results = _all_results_for_run(job_id)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["id","status","duration_ms","nl_query","run_sql_status","run_sql_rows",
                "run_sql_columns","generate_viz_status","chart_type",
                "fail_reasons","warn_reasons","error"])
    for r in results:
        v = r.get("validations") or {}
        rs = r.get("run_sql_call") or {}; gv = r.get("generate_viz_call") or {}
        w.writerow([
            r.get("id"), r.get("overall_status"), r.get("total_duration_ms"),
            (r.get("nl_query") or "")[:200], rs.get("status"),
            v.get("run_sql_row_count"),
            "|".join(map(str, v.get("run_sql_columns") or [])),
            gv.get("status"), v.get("generate_viz_chart_type"),
            "|".join(v.get("fail_reasons") or []),
            "|".join(v.get("warn_reasons") or []),
            (r.get("error") or "")[:300].replace("\n", " "),
        ])
    return Response(buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{job_id}-results.csv"'})


@app.route("/jobs/<job_id>/export/failures.xlsx")
@auth.login_required
def job_export_failures_xlsx(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    from openpyxl import Workbook
    results = _all_results_for_run(job_id)
    bad = [r for r in results if r.get("overall_status") in ("FAIL","PARTIAL","TIMEOUT")]
    wb = Workbook(); ws = wb.active; ws.title = "Failed queries"
    ws.append(["question","expected_sql","status","duration_ms",
               "fail_reasons","warn_reasons","run_sql_status","generate_viz_status","error_summary"])
    for r in bad:
        v = r.get("validations") or {}; rs = r.get("run_sql_call") or {}; gv = r.get("generate_viz_call") or {}
        err = (r.get("error") or "").splitlines()[0][:200] if r.get("error") else ""
        ws.append([
            r.get("nl_query"), r.get("expected_sql") or "",
            r.get("overall_status"), r.get("total_duration_ms"),
            "; ".join(v.get("fail_reasons") or []),
            "; ".join(v.get("warn_reasons") or []),
            rs.get("status"), gv.get("status"), err,
        ])
    ws.column_dimensions["A"].width = 60; ws.column_dimensions["B"].width = 40
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{job_id}-failures.xlsx"',
                 "X-Failure-Count": str(len(bad)), "X-Total-Count": str(len(results))})


@app.route("/jobs/<job_id>/download")
@auth.login_required
def job_download_zip(job_id: str):
    if not _run_owned_by_current_user(job_id): abort(404)
    job = _job_dir(job_id)
    if not job.exists(): abort(404)
    if not (job / "REPORT.html").exists():
        try: generate_report(job)
        except Exception: return ("Could not generate report yet", 500)
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for include in ("REPORT.html","REPORT.md","SUMMARY.json","all_results.json","job.json","questions.xlsx"):
            p = job / include
            if p.exists(): z.write(p, arcname=p.name)
        for sub in ("screenshots","results","network_logs"):
            d = job / sub
            if d.exists():
                for f in d.rglob("*"):
                    if f.is_file(): z.write(f, arcname=f"{sub}/{f.name}")
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{job_id}-report.zip"'})


# ---------------------------------------------------------------------------
# JSON APIs (all user-scoped)
# ---------------------------------------------------------------------------
def _runs_for_user_query(user_filter_clause: str = "") -> str:
    return f"""
    SELECT r.id, r.user_id, r.login_url, r.username, r.machine_id, r.sql_agent_path,
           r.question_count, r.status, r.summary, r.created_at, r.started_at, r.finished_at,
           r.test_file_name,
           r.assigned_to_name, r.assigned_to_id, r.assigned_at, r.minijira_task_id,
           COALESCE((SELECT COUNT(*) FROM query_results q WHERE q.run_id = r.id AND q.status = 'PASS'), 0)    AS pass_n,
           COALESCE((SELECT COUNT(*) FROM query_results q WHERE q.run_id = r.id AND q.status = 'PARTIAL'), 0) AS partial_n,
           COALESCE((SELECT COUNT(*) FROM query_results q WHERE q.run_id = r.id AND q.status = 'FAIL'), 0)    AS fail_n,
           COALESCE((SELECT COUNT(*) FROM query_results q WHERE q.run_id = r.id AND q.status = 'TIMEOUT'), 0) AS timeout_n
      FROM runs r
     {user_filter_clause}
     ORDER BY r.created_at DESC
    """


def _list_user_runs(limit: int | None = None) -> list[dict[str, Any]]:
    sql = _runs_for_user_query("WHERE r.user_id = %s")
    params: tuple = (auth.current_user_id(),)
    if limit: sql += f" LIMIT {int(limit)}"
    rows = db.fetch_all(sql, params)
    out = []
    for r in rows:
        out.append({
            "id": r["id"], "login_url": r["login_url"], "username": r["username"],
            "machine_id": r["machine_id"], "sql_agent_path": r["sql_agent_path"],
            "question_count": r["question_count"], "status": r["status"],
            "summary": {"PASS": r["pass_n"], "PARTIAL": r["partial_n"],
                        "FAIL": r["fail_n"], "TIMEOUT": r["timeout_n"]},
            "created_at": r["created_at"].isoformat() if r["created_at"] else None,
            "started_at": r["started_at"].isoformat() if r["started_at"] else None,
            "finished_at": r["finished_at"].isoformat() if r["finished_at"] else None,
            "test_file_name": r.get("test_file_name"),
            "assigned_to_name":  r.get("assigned_to_name"),
            "assigned_to_id":    r.get("assigned_to_id"),
            "assigned_at":       r["assigned_at"].isoformat() if r.get("assigned_at") else None,
            "minijira_task_id":  r.get("minijira_task_id"),
        })
    return out


@app.route("/api/runs")
@auth.login_required
def api_runs():
    runs = _list_user_runs()
    return jsonify({"runs": runs, "count": len(runs)})


@app.route("/api/runs/export.csv")
@auth.login_required
def api_runs_csv():
    runs = _list_user_runs()
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["id","created_at","started_at","finished_at","status",
                "login_url","username","questions","PASS","PARTIAL","FAIL","TIMEOUT"])
    for r in runs:
        s = r.get("summary") or {}
        w.writerow([r.get("id"), r.get("created_at"), r.get("started_at"), r.get("finished_at"),
                    r.get("status"), r.get("login_url"), r.get("username"),
                    r.get("question_count") or 0,
                    s.get("PASS",0), s.get("PARTIAL",0), s.get("FAIL",0), s.get("TIMEOUT",0)])
    return Response(buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": 'attachment; filename="all-runs.csv"'})


@app.route("/api/stats")
@auth.login_required
def api_stats():
    runs = _list_user_runs()
    now = datetime.now(timezone.utc)
    cutoff_7d = now - timedelta(days=7)
    runs_last_7d = queries_last_7d = 0
    total_pass = total_partial = total_fail = total_timeout = 0
    durations: list[int] = []; failure_counter: Counter[str] = Counter()
    daily: dict[str, dict[str, int]] = {}

    for r in runs:
        try:
            created_dt = datetime.fromisoformat(r["created_at"]) if r.get("created_at") else None
        except Exception:
            created_dt = None
        s = r.get("summary") or {}
        run_total = sum(s.get(k, 0) for k in ("PASS","PARTIAL","FAIL","TIMEOUT"))
        if created_dt and created_dt >= cutoff_7d:
            runs_last_7d += 1; queries_last_7d += run_total
        total_pass += s.get("PASS", 0); total_partial += s.get("PARTIAL", 0)
        total_fail += s.get("FAIL", 0); total_timeout += s.get("TIMEOUT", 0)
        if created_dt:
            day = created_dt.strftime("%Y-%m-%d")
            d = daily.setdefault(day, {"queries": 0, "pass": 0, "fail": 0})
            d["queries"] += run_total
            d["pass"] += s.get("PASS", 0)
            d["fail"] += s.get("FAIL", 0) + s.get("TIMEOUT", 0)
        for qr in _all_results_for_run(r["id"]):
            if qr.get("total_duration_ms"): durations.append(qr["total_duration_ms"])
            v = qr.get("validations") or {}
            for fr in v.get("fail_reasons") or []:
                failure_counter[fr] += 1

    total_queries = total_pass + total_partial + total_fail + total_timeout
    avg_ms = int(sum(durations) / len(durations)) if durations else 0
    series = []
    for i in range(29, -1, -1):
        day = (now - timedelta(days=i)).strftime("%Y-%m-%d")
        d = daily.get(day, {"queries": 0, "pass": 0, "fail": 0})
        series.append({"date": day[5:], **d})
    return jsonify({
        "totals": {"runs": len(runs), "queries": total_queries,
                   "pass": total_pass, "partial": total_partial,
                   "fail": total_fail, "timeout": total_timeout, "avg_ms": avg_ms,
                   "runs_last_7d": runs_last_7d, "queries_last_7d": queries_last_7d},
        "daily": series,
        "top_failures": [{"reason": r, "count": n} for r, n in failure_counter.most_common(10)],
        "recent_runs": [{"id": r["id"], "status": r["status"],
                         "created_at": r["created_at"], "summary": r["summary"]}
                        for r in runs[:10]],
    })


@app.route("/api/preview-questions", methods=["POST"])
@auth.login_required
def api_preview_questions():
    f = request.files.get("questions")
    if not f: return jsonify({"error": "no file uploaded"}), 400
    tmp = DATA_DIR / f"_preview_{uuid.uuid4().hex[:8]}.xlsx"
    try:
        f.save(str(tmp))
        try: qs = read_questions(tmp)
        except Exception as e: return jsonify({"error": str(e)}), 400
        warnings: list[str] = []
        for q in qs[:20]:
            txt = (q.get("natural_language_query") or "").strip()
            if len(txt) < 5:
                warnings.append(f"Row {q['id']}: question is very short — \"{txt}\"")
            elif txt[0].isdigit() and ":" in txt[:8]:
                warnings.append(f"Row {q['id']}: looks like a timestamp — \"{txt[:30]}\". Make sure column A contains the question, not a date.")
        return jsonify({"total": len(qs), "preview": qs[:20], "warnings": list(dict.fromkeys(warnings))[:10]})
    finally:
        try: tmp.unlink()
        except Exception: pass


# Test files (per-user)
@app.route("/api/test-files", methods=["GET", "POST"])
@auth.login_required
def api_test_files():
    if request.method == "GET":
        rows = db.fetch_all(
            "SELECT * FROM test_files WHERE user_id = %s ORDER BY uploaded_at DESC",
            (auth.current_user_id(),),
        )
        return jsonify({"files": [
            {"id": r["id"], "filename": r["filename"], "original_name": r["original_name"],
             "question_count": r["question_count"], "size_bytes": r["size_bytes"],
             "uploaded_at": r["uploaded_at"].isoformat() if r["uploaded_at"] else None}
            for r in rows]})
    f = request.files.get("file")
    if not f: return jsonify({"error": "no file uploaded (field name: 'file')"}), 400
    file_id = uuid.uuid4().hex[:10]
    safe = (f.filename or "upload.xlsx").replace("/", "_").replace("\\", "_")
    if not safe.lower().endswith(".xlsx"): safe += ".xlsx"
    stored_name = f"{file_id}-{safe}"
    stored_path = TEST_FILES_DIR / stored_name
    f.save(str(stored_path))
    try: qs = read_questions(stored_path)
    except Exception as e:
        try: stored_path.unlink()
        except Exception: pass
        return jsonify({"error": f"Could not parse uploaded .xlsx: {e}"}), 400
    db.execute(
        """INSERT INTO test_files (id, user_id, filename, original_name, question_count, size_bytes)
           VALUES (%s, %s, %s, %s, %s, %s)""",
        (file_id, auth.current_user_id(), stored_name, safe, len(qs), stored_path.stat().st_size),
    )
    row = db.fetch_one("SELECT * FROM test_files WHERE id = %s", (file_id,))
    return jsonify({"file": {
        "id": row["id"], "filename": row["filename"], "original_name": row["original_name"],
        "question_count": row["question_count"], "size_bytes": row["size_bytes"],
        "uploaded_at": row["uploaded_at"].isoformat() if row["uploaded_at"] else None,
    }}), 201


def _user_test_file(file_id: str) -> dict[str, Any] | None:
    return db.fetch_one(
        "SELECT * FROM test_files WHERE id = %s AND user_id = %s",
        (file_id, auth.current_user_id()),
    )


@app.route("/api/test-files/<file_id>", methods=["DELETE"])
@auth.login_required
def api_test_file_delete(file_id: str):
    row = _user_test_file(file_id)
    if not row: return jsonify({"error": "not found"}), 404
    try: (TEST_FILES_DIR / row["filename"]).unlink()
    except Exception: pass
    db.execute("DELETE FROM test_files WHERE id = %s", (file_id,))
    return jsonify({"ok": True})


@app.route("/api/test-files/<file_id>/preview")
@auth.login_required
def api_test_file_preview(file_id: str):
    row = _user_test_file(file_id)
    if not row: return jsonify({"error": "not found"}), 404
    try: qs = read_questions(TEST_FILES_DIR / row["filename"])
    except Exception as e: return jsonify({"error": str(e)}), 500
    return jsonify({"total": len(qs), "preview": qs[:20], "meta": {
        "id": row["id"], "original_name": row["original_name"],
        "uploaded_at": row["uploaded_at"].isoformat() if row["uploaded_at"] else None,
    }})


@app.route("/api/test-files/<file_id>/download")
@auth.login_required
def api_test_file_download(file_id: str):
    row = _user_test_file(file_id)
    if not row: abort(404)
    return send_from_directory(TEST_FILES_DIR, row["filename"],
        as_attachment=True, download_name=row["original_name"])


# Presets (per-user)
@app.route("/api/presets", methods=["GET", "POST"])
@auth.login_required
def api_presets():
    if request.method == "GET":
        rows = db.fetch_all(
            "SELECT * FROM presets WHERE user_id = %s ORDER BY created_at DESC",
            (auth.current_user_id(),))
        return jsonify({"presets": [
            {"id": r["id"], "name": r["name"], "login_url": r["login_url"],
             "username": r["username"], "machine_id": r["machine_id"],
             "sql_agent_path": r["sql_agent_path"],
             "run_sql_timeout_ms": r["run_sql_timeout_ms"],
             "gen_viz_timeout_ms": r["gen_viz_timeout_ms"],
             "created_at": r["created_at"].isoformat() if r["created_at"] else None}
            for r in rows]})
    body = request.get_json(silent=True) or {}
    if not body.get("name") or not body.get("login_url"):
        return jsonify({"error": "name and login_url are required"}), 400
    pid = uuid.uuid4().hex[:10]
    db.execute(
        """INSERT INTO presets (id, user_id, name, login_url, username,
                                 machine_id, sql_agent_path,
                                 run_sql_timeout_ms, gen_viz_timeout_ms)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (pid, auth.current_user_id(),
         str(body["name"]).strip()[:80], str(body["login_url"]).strip(),
         str(body.get("username") or "").strip(),
         str(body.get("machine_id") or "100").strip(),
         str(body.get("sql_agent_path") or "/backoffice/mv-assets/index-modern.html#/listScreen/sqlagent").strip(),
         int(body.get("run_sql_timeout_ms") or 120_000),
         int(body.get("gen_viz_timeout_ms") or 120_000)),
    )
    row = db.fetch_one("SELECT * FROM presets WHERE id = %s", (pid,))
    return jsonify({"preset": {k: (v.isoformat() if hasattr(v, "isoformat") else v)
                               for k, v in row.items()}}), 201


@app.route("/api/presets/<preset_id>", methods=["DELETE"])
@auth.login_required
def api_preset_delete(preset_id: str):
    row = db.fetch_one("SELECT id FROM presets WHERE id = %s AND user_id = %s",
                       (preset_id, auth.current_user_id()))
    if not row: return jsonify({"error": "not found"}), 404
    db.execute("DELETE FROM presets WHERE id = %s", (preset_id,))
    return jsonify({"ok": True})


# Whoami (handy for the UI + debugging)
@app.route("/api/me")
@auth.login_required
def api_me():
    u = g.user.copy()
    if u.get("created_at"): u["created_at"] = u["created_at"].isoformat()
    if u.get("last_login_at"): u["last_login_at"] = u["last_login_at"].isoformat()
    return jsonify({"user": u})


# ---------------------------------------------------------------------------
def main() -> None:
    import argparse, webbrowser
    ap = argparse.ArgumentParser()
    ap.add_argument("--host", default="127.0.0.1")
    ap.add_argument("--port", type=int, default=5050)
    ap.add_argument("--no-browser", action="store_true")
    args = ap.parse_args()
    url = f"http://{args.host}:{args.port}/"
    print(f"Skylar IQ QA Tool v{APP_VERSION} — {url}")
    if not args.no_browser:
        try: threading.Timer(1.0, lambda: webbrowser.open(url)).start()
        except Exception: pass
    app.run(host=args.host, port=args.port, debug=False, threaded=True)


if __name__ == "__main__":
    main()
